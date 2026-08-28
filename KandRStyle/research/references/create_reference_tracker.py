#!/usr/bin/env python3
"""
Script to create a Reference Tracking workbook for URLs in ut.tex
"""

import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
from pathlib import Path
from datetime import datetime

# Get the script directory and construct relative paths
SCRIPT_DIR = Path(__file__).parent.resolve()
# ut.tex is at KandRStyle/ut.tex, this script is at KandRStyle/research/references/
UT_TEX_PATH = SCRIPT_DIR / ".." / ".." / "ut.tex"
WORKBOOK_PATH = SCRIPT_DIR / "Reference_Tracking.xlsx"

def extract_urls_from_tex(file_path):
    """
    Extract URLs and their line numbers from the tex file
    Returns: List of tuples (line_number, url)
    """
    url_pattern = r'https?://[^\s}\]"\\]+'
    urls_with_lines = []
    
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()
        
    for line_num, line in enumerate(lines, 1):
        # Find all URLs in the line
        matches = re.finditer(url_pattern, line)
        for match in matches:
            url = match.group(0)
            # Clean up the URL (remove trailing punctuation that's not part of the URL)
            url = url.rstrip('.,;:')
            urls_with_lines.append((line_num, url))
    
    return urls_with_lines

def extract_accessed_dates_from_tex(file_path):
    """
    Extract URLs and their "Accessed" dates from ut.tex
    Returns: dict mapping URLs to accessed date strings
    """
    url_dates = {}
    
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        content = f.read()
    
    # Pattern to find endnotes with URLs and Accessed dates
    # Example: \url{https://www.mybib.com}. Accessed 12 July 2026.
    pattern = r'\\url\{([^}]+)\}[^}]*?Accessed\s+(\d+)\s+(\w+)\.?\s+(\d{4})'
    
    matches = re.finditer(pattern, content, re.IGNORECASE)
    
    for match in matches:
        url = match.group(1)
        day = match.group(2)
        month_name = match.group(3)
        year = match.group(4)
        
        # Store as formatted string for display
        date_str = f"{day} {month_name} {year}"
        
        # Try to parse to validate
        try:
            date_obj = datetime.strptime(date_str, "%d %B %Y")
            url_dates[url] = date_str
        except ValueError:
            try:
                # Try abbreviated month
                date_obj = datetime.strptime(date_str, "%d %b %Y")
                url_dates[url] = date_str
            except ValueError:
                # Skip if date can't be parsed, but store the raw string anyway
                url_dates[url] = date_str
    
    return url_dates

def create_reference_tracking_workbook(urls_with_lines, url_dates):
    """
    Create an Excel workbook with reference tracking information
    """
    # Create a new workbook
    wb = openpyxl.Workbook()
    
    # Remove the default sheet
    wb.remove(wb.active)
    
    # Create "Reference Tracking" sheet
    ws_tracking = wb.create_sheet("Reference Tracking")
    
    # Define headers for Reference Tracking sheet
    headers_tracking = ["ID", "URL", "Accessed Date", "Status", "Notes"]
    ws_tracking.append(headers_tracking)
    
    # Style the header row
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in ws_tracking[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Get unique URLs (preserve order of first occurrence)
    seen_urls = {}
    unique_urls = []
    for line_num, url in urls_with_lines:
        if url not in seen_urls:
            seen_urls[url] = []
            unique_urls.append(url)
        seen_urls[url].append(line_num)
    
    # Add URLs to Reference Tracking sheet
    for idx, url in enumerate(unique_urls, 1):
        accessed_date = url_dates.get(url, "")  # Get accessed date if available
        ws_tracking.append([idx, url, accessed_date, "Not Found", ""])
    
    # Adjust column widths
    ws_tracking.column_dimensions['A'].width = 8
    ws_tracking.column_dimensions['B'].width = 80
    ws_tracking.column_dimensions['C'].width = 18
    ws_tracking.column_dimensions['D'].width = 15
    ws_tracking.column_dimensions['E'].width = 30
    
    # Create "URL Locations" sheet
    ws_locations = wb.create_sheet("URL Locations")
    
    # Define headers for URL Locations sheet
    headers_locations = ["ID", "URL", "Line Numbers in ut.tex"]
    ws_locations.append(headers_locations)
    
    # Style the header row
    for cell in ws_locations[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add URL location information
    for idx, url in enumerate(unique_urls, 1):
        line_numbers = seen_urls[url]
        line_numbers_str = ", ".join(map(str, line_numbers))
        ws_locations.append([idx, url, line_numbers_str])
    
    # Adjust column widths
    ws_locations.column_dimensions['A'].width = 8
    ws_locations.column_dimensions['B'].width = 80
    ws_locations.column_dimensions['C'].width = 30
    
    # Save the workbook
    wb.save(str(WORKBOOK_PATH))
    
    print(f"Created workbook: {WORKBOOK_PATH.resolve()}")
    print(f"Total unique URLs found: {len(unique_urls)}")
    print(f"Total URL occurrences: {len(urls_with_lines)}")
    
    return unique_urls

def main():
    print(f"Extracting URLs from ut.tex at: {UT_TEX_PATH.resolve()}")
    urls_with_lines = extract_urls_from_tex(str(UT_TEX_PATH))
    
    if not urls_with_lines:
        print("No URLs found in ut.tex")
        return
    
    print(f"Found {len(urls_with_lines)} URL occurrences")
    
    print("\nExtracting 'Accessed' dates from ut.tex...")
    url_dates = extract_accessed_dates_from_tex(str(UT_TEX_PATH))
    print(f"Found {len(url_dates)} URLs with 'Accessed' dates")
    
    print("\nCreating Reference Tracking workbook...")
    unique_urls = create_reference_tracking_workbook(urls_with_lines, url_dates)
    
    print("\nFirst 10 URLs to process:")
    for idx, url in enumerate(unique_urls[:10], 1):
        accessed_info = f" (Accessed: {url_dates[url]})" if url in url_dates else " (No accessed date)"
        print(f"{idx}. {url}{accessed_info}")

if __name__ == "__main__":
    main()
