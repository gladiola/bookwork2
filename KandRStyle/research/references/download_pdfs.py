#!/usr/bin/env python3
"""
Script to download first 10 URLs as PDFs using playwright
"""

import openpyxl
import os
import sys
import time
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

WORKBOOK_PATH = "/home/runner/work/bookwork2/bookwork2/KandRStyle/research/references/Reference_Tracking.xlsx"
PDF_DIR = "/home/runner/work/bookwork2/bookwork2/KandRStyle/research/references"

def load_urls_from_workbook():
    """Load URLs from the Reference Tracking workbook"""
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["Reference Tracking"]
    
    urls = []
    for row in ws.iter_rows(min_row=2, max_row=11, values_only=True):  # Get first 10 URLs (rows 2-11)
        if row[0] and row[1]:  # ID and URL columns
            urls.append({
                'id': row[0],
                'url': row[1],
                'status': row[2] if len(row) > 2 else "Not Found"
            })
    
    return urls

def update_workbook_status(url_id, status, notes=""):
    """Update the status of a URL in the workbook"""
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["Reference Tracking"]
    
    # Find the row with the matching ID
    for row in ws.iter_rows(min_row=2):
        if row[0].value == url_id:
            row[2].value = status  # Update Status column
            if notes:
                row[3].value = notes  # Update Notes column
            break
    
    wb.save(WORKBOOK_PATH)

def download_url_as_pdf(url, pdf_path, timeout=30000):
    """Download a URL as PDF using playwright"""
    try:
        with sync_playwright() as p:
            # Launch browser
            browser = p.chromium.launch(headless=True)
            context = browser.new_context()
            page = context.new_page()
            
            print(f"  Navigating to {url}...")
            # Navigate to the URL
            page.goto(url, timeout=timeout, wait_until="networkidle")
            
            # Wait a bit for any JavaScript to execute
            page.wait_for_timeout(2000)
            
            print(f"  Generating PDF...")
            # Generate PDF
            page.pdf(path=pdf_path, format='A4', print_background=True)
            
            browser.close()
            return True, "Successfully downloaded"
            
    except PlaywrightTimeoutError:
        return False, "Timeout - page took too long to load"
    except Exception as e:
        return False, f"Error: {str(e)[:100]}"

def main():
    print("Loading URLs from workbook...")
    urls = load_urls_from_workbook()
    
    print(f"\nFound {len(urls)} URLs to process\n")
    
    results = []
    
    for url_data in urls:
        url_id = url_data['id']
        url = url_data['url']
        pdf_filename = f"ref_{url_id:03d}.pdf"
        pdf_path = os.path.join(PDF_DIR, pdf_filename)
        
        print(f"Processing URL {url_id}: {url}")
        
        # Skip local IP addresses
        if "192.168." in url or "127.0.0.1" in url or "localhost" in url:
            print(f"  Skipping local/private URL")
            update_workbook_status(url_id, "Not Found", "Local/private URL - cannot access")
            results.append({
                'id': url_id,
                'url': url,
                'status': 'Skipped',
                'reason': 'Local/private URL'
            })
            continue
        
        # Try to download
        success, message = download_url_as_pdf(url, pdf_path)
        
        if success:
            print(f"  ✓ Success: {pdf_filename}")
            update_workbook_status(url_id, "Stored", f"Saved as {pdf_filename}")
            results.append({
                'id': url_id,
                'url': url,
                'status': 'Success',
                'filename': pdf_filename
            })
        else:
            print(f"  ✗ Failed: {message}")
            update_workbook_status(url_id, "Not Found", message)
            results.append({
                'id': url_id,
                'url': url,
                'status': 'Failed',
                'reason': message
            })
        
        # Small delay between requests
        time.sleep(2)
    
    # Print summary
    print("\n" + "="*80)
    print("SUMMARY")
    print("="*80)
    
    success_count = sum(1 for r in results if r['status'] == 'Success')
    failed_count = sum(1 for r in results if r['status'] == 'Failed')
    skipped_count = sum(1 for r in results if r['status'] == 'Skipped')
    
    print(f"\nTotal URLs processed: {len(results)}")
    print(f"Successfully downloaded: {success_count}")
    print(f"Failed: {failed_count}")
    print(f"Skipped: {skipped_count}")
    
    print("\nSuccessful downloads:")
    for r in results:
        if r['status'] == 'Success':
            print(f"  {r['id']}. {r['filename']} - {r['url'][:60]}...")
    
    if failed_count > 0:
        print("\nFailed downloads:")
        for r in results:
            if r['status'] == 'Failed':
                print(f"  {r['id']}. {r['url'][:60]}... - {r['reason']}")
    
    if skipped_count > 0:
        print("\nSkipped URLs:")
        for r in results:
            if r['status'] == 'Skipped':
                print(f"  {r['id']}. {r['url'][:60]}... - {r['reason']}")

if __name__ == "__main__":
    main()
