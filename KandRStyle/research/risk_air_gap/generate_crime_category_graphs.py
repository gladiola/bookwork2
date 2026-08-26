#!/usr/bin/env python3
from __future__ import annotations

import argparse
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
import openpyxl

YEARS = [2016, 2017, 2018, 2019, 2020, 2021, 2022]

ALIASES = {
    "ipr": "IPR/Counterfeit",
    "intellectual property rights": "IPR/Counterfeit",
    "intellectual property theft": "IPR/Counterfeit",
    "identity theft": "Identify Theft",
    "data breach": "Data Breach",
    "personal data breach": "Personal Data Breach",
    "threats": "Threats/Stalking/Harassment/Terrorism",
    "bec": "BEC",
}


def canonical_category(name: str) -> str:
    k = name.strip().lower()
    return ALIASES.get(k, name)


def display_category(category: str) -> str:
    if category == "IPR/Counterfeit":
        return "Intellectual Property Rights (IPR)"
    return category


def parse_num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("#"):
            return 0.0
        s = s.replace(",", "").replace("$", "")
        try:
            return float(s)
        except ValueError:
            return 0.0
    try:
        return float(v)
    except Exception:
        return 0.0


def fmt_num(v: float) -> str:
    if abs(v - round(v)) < 1e-9:
        return f"{int(round(v)):,}"
    return f"{v:,.2f}"


def fmt_dollars(v: float, decimals: int = 0) -> str:
    return f"\\${v:,.{decimals}f}"


def header_col(ws, category: str) -> int:
    for c in range(1, ws.max_column + 1):
        if ws.cell(4, c).value == category:
            return c
    raise ValueError(f"Category not found in summary sheet header: {category}")


def get_scaled_series(wb, category: str):
    ws_as = wb["City Summary AMERICAN_SAMOA"]
    ws_ca = wb["City Summary CALIFORNIA"]
    ws_ias = wb["City Summary Impact AMERICAN_SA"]
    ws_ica = wb["City Summary Impact CALIFORNIA"]

    col = header_col(ws_as, category)
    as_counts = [parse_num(ws_as.cell(r, col).value) for r in range(5, 12)]
    ca_counts = [parse_num(ws_ca.cell(r, col).value) for r in range(5, 12)]
    as_imp = [parse_num(ws_ias.cell(r, col).value) for r in range(5, 12)]
    ca_imp = [parse_num(ws_ica.cell(r, col).value) for r in range(5, 12)]
    return as_counts, ca_counts, as_imp, ca_imp


def get_raw_for_year(wb, category: str, year: int, region: str):
    if region == "AS":
        sname = f"{year}_IC3 AMERICAN_SAMOA"
    else:
        sname = f"{year}_IC3 State CALIFORNIA"
    ws = wb[sname]
    row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 3).value == category:
            row = r
            break
    if row is None:
        return 0.0, 0.0
    count = parse_num(ws.cell(row, 5).value)
    impact = parse_num(ws.cell(row, 8).value)
    return count, impact


def simulate_counts(lower_mean: float, upper_mean: float, n: int, rng: np.random.Generator):
    lo, hi = sorted((lower_mean, upper_mean))
    if lo >= 0 and hi <= 1:
        values = [0, 1]
        p = (lo + hi) / 2
        samples = rng.binomial(1, p, size=n)
        dist = "binary"
    else:
        low_i = int(math.floor(lo))
        high_i = int(math.ceil(hi))
        if high_i <= low_i:
            high_i = low_i + 1
        values = list(range(low_i, high_i + 1))
        probs = np.full(len(values), 1 / len(values))
        samples = rng.choice(values, size=n, p=probs)
        dist = "discrete"
    pct = [float(np.mean(samples == v) * 100.0) for v in values]
    return dist, values, pct, samples


def simulate_impacts(lower: float, upper: float, n: int, rng: np.random.Generator):
    lo, hi = sorted((lower, upper))
    if math.isclose(lo, hi):
        return np.full(n, lo)
    return rng.uniform(lo, hi, size=n)


def plot_count(values, pct, dist, category, slug, out_dir: Path):
    fig, ax = plt.subplots(figsize=(10, 6), facecolor="white")
    ax.set_facecolor("white")
    colors = ["white" if i % 2 == 0 else "black" for i in range(len(values))]
    bars = ax.bar(values, pct, color=colors, edgecolor="black", linewidth=1.5)
    for b, p in zip(bars, pct):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.3, f"{p:.1f}%", ha="center", va="bottom")
    ax.set_xlabel("Events")
    ax.set_ylabel("Likelihood (%)")
    ax.set_xticks(values)
    ax.yaxis.set_major_formatter(ticker.FormatStrFormatter("%.1f"))
    ax.set_title(f"Monte Carlo {dist.title()} Simulation: {category} Count (10,000 Iterations)")
    ax.grid(axis="y", linestyle="--", color="gray", alpha=0.4)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    plt.tight_layout()
    name = f"{dist}_{slug}_count.png"
    fig.savefig(out_dir / name, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return name


def plot_impact_hist(samples, category, slug, out_dir: Path):
    fig, ax = plt.subplots(figsize=(10, 6), facecolor="white")
    ax.set_facecolor("white")
    ax.hist(samples, bins=50, color="white", edgecolor="black", linewidth=0.8)
    ax.set_xlabel("Impact ($)")
    ax.set_ylabel("Count")
    ax.xaxis.set_major_formatter(ticker.FuncFormatter(lambda x, _: f"${x:,.0f}"))
    ax.set_title(f"Monte Carlo Simulation: {category} Impact (10,000 Iterations)")
    ax.grid(axis="y", linestyle="--", color="gray", alpha=0.4)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    plt.tight_layout()
    name = f"histogram_{slug}_impact_bw.png"
    fig.savefig(out_dir / name, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return name


def plot_impact_percentile(samples, category, slug, out_dir: Path):
    s = np.sort(samples)
    pcts = np.arange(0, 101)
    vals = np.percentile(s, pcts)
    fig, ax = plt.subplots(figsize=(10, 6), facecolor="white")
    ax.set_facecolor("white")
    ax.plot(vals, 100 - pcts, color="black", linewidth=2)
    for level, ls in [(90, "--"), (50, ":"), (10, "-.")]:
        ax.axhline(level, color="gray", linestyle=ls, linewidth=1)
    ax.set_xlabel("Cost ($)")
    ax.set_ylabel("Probability of Exceeding")
    ax.xaxis.set_major_formatter(ticker.FuncFormatter(lambda x, _: f"${x:,.0f}"))
    ax.yaxis.set_major_formatter(ticker.FormatStrFormatter("%d%%"))
    ax.set_title(f"Monte Carlo Simulation: {category} Impact Percentiles")
    ax.grid(linestyle="--", color="gray", alpha=0.35)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    plt.tight_layout()
    name = f"percent_{slug}_impact_bw.png"
    fig.savefig(out_dir / name, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    p10, p50, p90 = np.percentile(s, [10, 50, 90])
    return name, float(p10), float(p50), float(p90)


def render_table(values, pcts):
    top_vals = values[:5] + [""] * max(0, 5 - len(values[:5]))
    bot_vals = values[5:10] + [""] * max(0, 5 - len(values[5:10]))
    top_pct = [f"{v:.1f}\\%" for v in pcts[:5]] + [""] * max(0, 5 - len(pcts[:5]))
    bot_pct = [f"{v:.1f}\\%" for v in pcts[5:10]] + [""] * max(0, 5 - len(pcts[5:10]))
    return top_vals, top_pct, bot_vals, bot_pct


def build_text(category: str, slug: str, as_raw, ca_raw, as_scaled, ca_scaled, means, sim, figure_names):
    category_display = display_category(category)
    lower_m, upper_m, lower_d, upper_d = means
    dist, values, pcts = sim
    count_fig, impact_hist_fig, impact_pct_fig, p10, p50, p90 = figure_names
    separation = abs(upper_m - lower_m)
    top_vals, top_pct, bot_vals, bot_pct = render_table(values, pcts)

    def row(vals):
        return " & \\cdots & ".join(vals)

    txt = []
    txt.append(f"%%%%%%%%%%%%%%%%%%%%  {category_display.upper()} %%%%%%%%%%%%%%%%%%%%%%%%")
    txt.append("\\newpage")
    txt.append(f"\\subsection{{{category_display}}}")
    txt.append("\\subsubsection{Lower Bound Computation}")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\mathbf{x}_{\\text{American Samoa}}^{(\\text{count})}\\\\")
    txt.append("\\mathbf{x}_{\\text{American Samoa}}^{(\\text{dollars})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{row([fmt_num(as_raw[0][0]), fmt_num(as_raw[1][0]), fmt_num(as_raw[2][0])])}\\\\")
    txt.append(f"{row([fmt_dollars(as_raw[0][1]), fmt_dollars(as_raw[1][1]), fmt_dollars(as_raw[2][1])])}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\mathbf{x}_{\\text{American Samoa}}^{(\\text{count scaled to sample population})}\\\\")
    txt.append("\\mathbf{x}_{\\text{American Samoa}}^{(\\text{dollars scaled to sample population})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{row([fmt_num(as_scaled[0][0]), fmt_num(as_scaled[1][0]), fmt_num(as_scaled[2][0])])}\\\\")
    txt.append(f"{row([fmt_dollars(as_scaled[0][1]), fmt_dollars(as_scaled[1][1]), fmt_dollars(as_scaled[2][1])])}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\bar{X}_{\\text{American Samoa}}^{(\\text{scaled count})}\\\\")
    txt.append("\\bar{X}_{\\text{American Samoa}}^{(\\text{scaled dollars})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{lower_m:.2f}\\\\")
    txt.append(f"{fmt_dollars(lower_d)}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\subsubsection{Upper Bound Computation}")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\mathbf{x}_{\\text{California}}^{(\\text{count})}\\\\")
    txt.append("\\mathbf{x}_{\\text{California}}^{(\\text{dollars})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{row([fmt_num(ca_raw[0][0]), fmt_num(ca_raw[1][0]), fmt_num(ca_raw[2][0])])}\\\\")
    txt.append(f"{row([fmt_dollars(ca_raw[0][1]), fmt_dollars(ca_raw[1][1]), fmt_dollars(ca_raw[2][1])])}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\mathbf{x}_{\\text{California}}^{(\\text{count scaled to sample population})}\\\\")
    txt.append("\\mathbf{x}_{\\text{California}}^{(\\text{dollars scaled to sample population})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{row([fmt_num(ca_scaled[0][0]), fmt_num(ca_scaled[1][0]), fmt_num(ca_scaled[2][0])])}\\\\")
    txt.append(f"{row([fmt_dollars(ca_scaled[0][1]), fmt_dollars(ca_scaled[1][1]), fmt_dollars(ca_scaled[2][1])])}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\[")
    txt.append("\\begin{bmatrix}")
    txt.append("\\bar{X}_{\\text{California}}^{(\\text{scaled count})}\\\\")
    txt.append("\\bar{X}_{\\text{California}}^{(\\text{scaled dollars})}")
    txt.append("\\end{bmatrix}")
    txt.append("=")
    txt.append("\\begin{bmatrix}")
    txt.append(f"{upper_m:.2f}\\\\")
    txt.append(f"{fmt_dollars(upper_d)}")
    txt.append("\\end{bmatrix}.")
    txt.append("\\]")
    txt.append("\\subsubsection{Monte Carlo Simulations for Likelihood and Impact}")
    txt.append("\\begin{figure}[H]")
    txt.append("  \\centering")
    txt.append(f"  \\includegraphics[scale=0.25]{{figures/{count_fig}}}")
    txt.append(f"  \\caption{{We ran a Monte Carlo simulation of 10,000 iterations with a {dist} distribution and bounded values between {min(values)} and {max(values)} events.}}")
    txt.append("\\end{figure}")
    txt.append(f"We took the mean of the sample from American Samoa for the lower bound ({lower_m:.2f}); we took the mean from the sample from California for the upper bound ({upper_m:.2f}); we had a separation of {separation:.2f} units and {len(values)} categories in-range.")
    txt.append("\\begin{table}[H]")
    txt.append("\\centering")
    txt.append(f"\\caption{{Monte Carlo {'Binary' if dist=='binary' else 'Discrete'} Simulation Values for {category_display}}}")
    txt.append("\\begin{tabular}{@{}>{\\centering\\arraybackslash}p{0.20\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}}")
    txt.append("\\toprule")
    txt.append(f"Events &{top_vals[0]} &{top_vals[1]} &{top_vals[2]} &{top_vals[3]} &{top_vals[4]}\\\\")
    txt.append("\\midrule")
    txt.append(f"Likelihood &{top_pct[0]} &{top_pct[1]} &{top_pct[2]} &{top_pct[3]} &{top_pct[4]}\\\\")
    txt.append("\\bottomrule")
    txt.append("\\end{tabular}")
    txt.append("\\begin{tabular}{@{}>{\\centering\\arraybackslash}p{0.20\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}>{\\centering\\arraybackslash}p{0.10\\linewidth}}")
    txt.append("\\toprule")
    txt.append(f"Events &{bot_vals[0]} &{bot_vals[1]} &{bot_vals[2]} &{bot_vals[3]} &{bot_vals[4]}\\\\")
    txt.append("\\midrule")
    txt.append(f"Likelihood &{bot_pct[0]} &{bot_pct[1]} &{bot_pct[2]} &{bot_pct[3]} &{bot_pct[4]}\\\\")
    txt.append("\\bottomrule")
    txt.append("\\end{tabular}")
    txt.append("\\end{table}")
    txt.append("\\begin{figure}[H]")
    txt.append("  \\centering")
    txt.append(f"  \\includegraphics[width=\\linewidth]{{figures/{impact_hist_fig}}}")
    txt.append(f"  \\caption{{Impact Monte Carlo histogram bounded between {fmt_dollars(lower_d)} and {fmt_dollars(upper_d)}.}}")
    txt.append("\\end{figure}")
    txt.append("\\begin{figure}[H]")
    txt.append("  \\centering")
    txt.append(f"  \\includegraphics[width=\\linewidth]{{figures/{impact_pct_fig}}}")
    txt.append("  \\[")
    txt.append(f"  \\mathrm{{CI}}_{{80\\%}}(\\theta) = [{fmt_dollars(p10)}, {fmt_dollars(p90)}]")
    txt.append("  \\]")
    txt.append(f"  \\caption{{There is a 90\\% probability costs are {fmt_dollars(p10,2)} or more, a 50\\% probability costs are {fmt_dollars(p50,2)} or more, and a 10\\% probability costs are {fmt_dollars(p90,2)} or more.}}")
    txt.append("\\end{figure}")
    return "\n".join(txt) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Generate black-and-white Monte Carlo graphs for any crime category.")
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--category", required=True)
    ap.add_argument("--figures-dir", required=True)
    ap.add_argument("--output-text", required=True)
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--iterations", type=int, default=10_000)
    args = ap.parse_args()

    category = canonical_category(args.category)
    slug = category.lower().replace("/", "").replace(" ", "").replace("-", "")
    figures_dir = Path(args.figures_dir)
    figures_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    as_counts, ca_counts, as_imp, ca_imp = get_scaled_series(wb, category)

    as_raw = [get_raw_for_year(wb, category, y, "AS") for y in (2016, 2017, 2022)]
    ca_raw = [get_raw_for_year(wb, category, y, "CA") for y in (2016, 2017, 2022)]

    as_scaled = [
        (as_counts[0], as_imp[0]),
        (as_counts[1], as_imp[1]),
        (as_counts[6], as_imp[6]),
    ]
    ca_scaled = [
        (ca_counts[0], ca_imp[0]),
        (ca_counts[1], ca_imp[1]),
        (ca_counts[6], ca_imp[6]),
    ]

    lower_count_mean = float(np.mean(as_counts))
    upper_count_mean = float(np.mean(ca_counts))
    lower_impact_mean = float(np.mean(as_imp))
    upper_impact_mean = float(np.mean(ca_imp))

    rng = np.random.default_rng(args.seed)
    dist, values, pcts, _ = simulate_counts(lower_count_mean, upper_count_mean, args.iterations, rng)
    impact_samples = simulate_impacts(lower_impact_mean, upper_impact_mean, args.iterations, rng)

    count_fig = plot_count(values, pcts, dist, category, slug, figures_dir)
    impact_hist_fig = plot_impact_hist(impact_samples, category, slug, figures_dir)
    impact_pct_fig, p10, p50, p90 = plot_impact_percentile(impact_samples, category, slug, figures_dir)

    text = build_text(
        category,
        slug,
        as_raw,
        ca_raw,
        as_scaled,
        ca_scaled,
        (lower_count_mean, upper_count_mean, lower_impact_mean, upper_impact_mean),
        (dist, values, pcts),
        (count_fig, impact_hist_fig, impact_pct_fig, p10, p50, p90),
    )
    out_text = Path(args.output_text)
    out_text.parent.mkdir(parents=True, exist_ok=True)
    out_text.write_text(text)

    print(f"Category: {category}")
    print(f"Distribution: {dist}")
    print(f"Lower mean count (AS): {lower_count_mean:.4f}")
    print(f"Upper mean count (CA): {upper_count_mean:.4f}")
    print(f"Lower mean impact (AS): {lower_impact_mean:.2f}")
    print(f"Upper mean impact (CA): {upper_impact_mean:.2f}")
    print(f"Count likelihoods: {dict(zip(values, [round(v, 2) for v in pcts]))}")
    print(f"Wrote: {out_text}")
    print(f"Figures: {count_fig}, {impact_hist_fig}, {impact_pct_fig}")


if __name__ == "__main__":
    main()
