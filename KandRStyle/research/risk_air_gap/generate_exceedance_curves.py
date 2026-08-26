#!/usr/bin/env python3
"""
Generate fresh exceedance and impact curves for Phishing, Threats, ID Theft, and Business Email Compromise.
"""

import openpyxl
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

# Use grayscale style
plt.style.use('grayscale')

def extract_crime_data(wb, crime_name, territory_name, years):
    """Extract crime data from worksheets."""
    data = {
        'raw_counts': [],
        'raw_impacts': [],
        'scaled_counts': [],
        'scaled_impacts': [],
        'years': years
    }
    
    for year in years:
        if 'AMERICAN_SAMOA' in territory_name.upper():
            sheet_name = f'{year}_IC3 AMERICAN_SAMOA'
        elif 'CALIFORNIA' in territory_name.upper():
            sheet_name = f'{year}_IC3 State CALIFORNIA'
        else:
            continue
            
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        
        # Find crime row
        for row_idx in range(1, min(100, ws.max_row + 1)):
            cell = ws.cell(row=row_idx, column=3)
            if cell.value and crime_name.lower() in str(cell.value).lower():
                raw_count = ws.cell(row=row_idx, column=5).value or 0
                scaled_count = ws.cell(row=row_idx, column=7).value or 0
                raw_impact = ws.cell(row=row_idx, column=8).value or 0
                scaled_impact = ws.cell(row=row_idx, column=10).value or 0
                
                # Handle error values
                if isinstance(scaled_count, str) and '#' in scaled_count:
                    scaled_count = 0
                if isinstance(scaled_impact, str) and '#' in scaled_impact:
                    scaled_impact = 0
                    
                data['raw_counts'].append(raw_count)
                data['raw_impacts'].append(raw_impact)
                data['scaled_counts'].append(scaled_count)
                data['scaled_impacts'].append(scaled_impact)
                break
    
    return data

def calculate_means(data):
    """Calculate mean of scaled counts and impacts."""
    scaled_counts = [v for v in data['scaled_counts'] if v != 0 and not isinstance(v, str)]
    scaled_impacts = [v for v in data['scaled_impacts'] if v != 0 and not isinstance(v, str)]
    
    mean_count = np.mean(scaled_counts) if scaled_counts else 0
    mean_impact = np.mean(scaled_impacts) if scaled_impacts else 0
    
    return {
        'mean_count': mean_count,
        'mean_impact': mean_impact,
    }

def generate_impact_curves(crime_name, as_mean, ca_mean, output_prefix):
    """Generate histogram and percentile curves for impact."""
    
    # Generate normal distribution for impact
    np.random.seed(42)
    mean = (as_mean + ca_mean) / 2
    std = (ca_mean - as_mean) / 4  # Approximate std dev
    
    if std <= 0:
        std = mean * 0.25 if mean > 0 else 100
    
    # Generate 10,000 samples
    samples = np.random.normal(mean, std, 10000)
    # Clip to reasonable bounds
    samples = np.clip(samples, max(0, as_mean * 0.5), ca_mean * 1.5)
    
    # 1. Create histogram
    fig, ax = plt.subplots(figsize=(12, 6))
    n, bins, patches = ax.hist(samples, bins=50, color='gray', edgecolor='black', linewidth=1)
    
    ax.set_xlabel('Impact Cost ($)', fontsize=12, fontweight='bold')
    ax.set_ylabel('Frequency', fontsize=12, fontweight='bold')
    ax.set_title(f'Monte Carlo Normal Distribution: {crime_name} Impact', fontsize=14, fontweight='bold')
    ax.grid(axis='y', alpha=0.3, linestyle='--')
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f'${x:,.0f}'))
    
    plt.tight_layout()
    plt.savefig(f'{output_prefix}_impact.png', dpi=150, bbox_inches='tight', facecolor='white')
    print(f'Saved {crime_name} impact histogram')
    plt.close()
    
    # 2. Create percentile/exceedance curve
    sorted_samples = np.sort(samples)
    percentiles = np.arange(0, 100, 0.1)
    values = np.percentile(sorted_samples, percentiles)
    exceedance = 100 - percentiles
    
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.plot(values, exceedance, color='black', linewidth=2)
    
    # Mark key percentiles
    p10 = np.percentile(sorted_samples, 10)
    p50 = np.percentile(sorted_samples, 50)
    p90 = np.percentile(sorted_samples, 90)
    
    ax.axvline(p10, color='gray', linestyle='--', linewidth=1, alpha=0.7)
    ax.axvline(p50, color='gray', linestyle='--', linewidth=1, alpha=0.7)
    ax.axvline(p90, color='gray', linestyle='--', linewidth=1, alpha=0.7)
    
    # Add labels
    y_pos = 95
    ax.text(p10, y_pos, f'P10: ${p10:,.0f}', rotation=0, fontsize=9)
    ax.text(p50, y_pos, f'P50: ${p50:,.0f}', rotation=0, fontsize=9)
    ax.text(p90, y_pos, f'P90: ${p90:,.0f}', rotation=0, fontsize=9)
    
    ax.set_xlabel('Impact Cost ($)', fontsize=12, fontweight='bold')
    ax.set_ylabel('Probability of Exceedance (%)', fontsize=12, fontweight='bold')
    ax.set_title(f'Exceedance Curve: {crime_name} Impact', fontsize=14, fontweight='bold')
    ax.grid(alpha=0.3, linestyle='--')
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f'${x:,.0f}'))
    
    plt.tight_layout()
    plt.savefig(f'{output_prefix}_impact.png'.replace('histogram', 'percent'), dpi=150, bbox_inches='tight', facecolor='white')
    print(f'Saved {crime_name} exceedance curve')
    plt.close()
    
    return {
        'p10': p10,
        'p50': p50,
        'p90': p90
    }

def main():
    print('Loading workbook...')
    wb = load_workbook('ComputerCrime_AUG2026_Impact_Final_With_C5C7_IFERROR_CHARTS_FIXED.xlsx', data_only=True)
    
    years = ['2016', '2017', '2018', '2019', '2020', '2021', '2022']
    
    # Crime categories to process
    crimes = [
        ('Phishing', 'phishing'),
        ('Threats', 'threats'),
        ('Identity Theft', 'idtheft'),
        ('BEC', 'bec')
    ]
    
    results = {}
    
    for crime_display, crime_file in crimes:
        print(f'\n=== Processing {crime_display} ===')
        
        # Extract data
        as_data = extract_crime_data(wb, crime_display, 'AMERICAN_SAMOA', years)
        ca_data = extract_crime_data(wb, crime_display, 'CALIFORNIA', years)
        
        # Calculate means
        as_means = calculate_means(as_data)
        ca_means = calculate_means(ca_data)
        
        print(f'  American Samoa mean impact: ${as_means["mean_impact"]:,.2f}')
        print(f'  California mean impact: ${ca_means["mean_impact"]:,.2f}')
        
        # Generate curves
        if as_means["mean_impact"] > 0 or ca_means["mean_impact"] > 0:
            output_prefix = f'../../figures/histogram_{crime_file}'
            percentiles = generate_impact_curves(
                crime_display,
                as_means['mean_impact'],
                ca_means['mean_impact'],
                output_prefix
            )
            
            results[crime_file] = {
                'as_mean': as_means['mean_impact'],
                'ca_mean': ca_means['mean_impact'],
                'percentiles': percentiles
            }
        else:
            print(f'  Skipping {crime_display} - insufficient data')
    
    print('\n=== Summary ===')
    for crime, data in results.items():
        print(f'\n{crime.upper()}:')
        print(f'  P10: ${data["percentiles"]["p10"]:,.2f}')
        print(f'  P50: ${data["percentiles"]["p50"]:,.2f}')
        print(f'  P90: ${data["percentiles"]["p90"]:,.2f}')

if __name__ == '__main__':
    main()
