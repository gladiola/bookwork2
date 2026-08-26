#!/usr/bin/env python3
"""
Process Extortion crime data for Monte Carlo risk assessment.
Extracts data from Excel, calculates bounds, and prepares output for LaTeX.
"""

import openpyxl
from openpyxl import load_workbook
import numpy as np
import sys
import json

def extract_extortion_data(wb, territory_name, years):
    """Extract Extortion data from worksheets for a given territory."""
    data = {
        'raw_counts': [],
        'raw_impacts': [],
        'scaled_counts': [],
        'scaled_impacts': [],
        'years': years
    }
    
    for year in years:
        if 'AMERICAN_SAMOA' in territory_name.upper():
            sheet_name = f'{year}_IC3 AMERICAN_SAMOA'
        elif 'CALIFORNIA' in territory_name.upper():
            sheet_name = f'{year}_IC3 State CALIFORNIA'
        else:
            continue
            
        if sheet_name not in wb.sheetnames:
            print(f'Warning: Sheet {sheet_name} not found', file=sys.stderr)
            continue
            
        ws = wb[sheet_name]
        
        # Find Extortion row
        for row_idx in range(1, min(100, ws.max_row + 1)):
            cell = ws.cell(row=row_idx, column=3)
            if cell.value and 'Extortion' in str(cell.value):
                # Column structure (0-indexed after column 3):
                # 0: Crime name
                # 1: "State Victim Count"
                # 2: Raw count
                # 3: Per capita rate
                # 4: Scaled count (to Chattanooga population)
                # 5: Raw dollar impact
                # 6: Per capita dollar impact
                # 7: Scaled dollar impact
                
                raw_count = ws.cell(row=row_idx, column=5).value or 0
                scaled_count = ws.cell(row=row_idx, column=7).value or 0
                raw_impact = ws.cell(row=row_idx, column=8).value or 0
                scaled_impact = ws.cell(row=row_idx, column=10).value or 0
                
                # Handle error values
                if isinstance(scaled_count, str) and '#' in scaled_count:
                    scaled_count = 0
                if isinstance(scaled_impact, str) and '#' in scaled_impact:
                    scaled_impact = 0
                    
                data['raw_counts'].append(raw_count)
                data['raw_impacts'].append(raw_impact)
                data['scaled_counts'].append(scaled_count)
                data['scaled_impacts'].append(scaled_impact)
                break
    
    return data

def calculate_means(data):
    """Calculate mean of scaled counts and impacts."""
    scaled_counts = [v for v in data['scaled_counts'] if v != 0 and not isinstance(v, str)]
    scaled_impacts = [v for v in data['scaled_impacts'] if v != 0 and not isinstance(v, str)]
    
    mean_count = np.mean(scaled_counts) if scaled_counts else 0
    mean_impact = np.mean(scaled_impacts) if scaled_impacts else 0
    
    return {
        'mean_count': mean_count,
        'mean_impact': mean_impact,
        'count_values': scaled_counts,
        'impact_values': scaled_impacts
    }

def determine_distribution_type(lower_bound, upper_bound):
    """
    Determine distribution type based on bounds.
    Returns: 'binary', 'discrete', or 'normal'
    """
    # Round to integers for count data
    lower = round(lower_bound)
    upper = round(upper_bound)
    
    # If both bounds are between 0 and 1, use binary
    if lower_bound <= 1 and upper_bound <= 1:
        return 'binary', lower, upper
    
    # Calculate the range
    range_size = upper - lower + 1
    
    # If range is 2 to 10 values, use discrete
    if 2 <= range_size <= 10:
        return 'discrete', lower, upper
    
    # If 25 or more, use normal
    if range_size >= 25:
        return 'normal', lower, upper
    
    # For 11-24, still use discrete
    return 'discrete', lower, upper

def format_number_with_commas(num):
    """Format number with thousands separators for LaTeX."""
    if num == 0:
        return '0'
    if isinstance(num, int) or num == int(num):
        return f'{int(num):,}'.replace(',', '{,}')
    else:
        # For floats, format with 2 decimal places
        return f'{num:,.2f}'.replace(',', '{,}')

def format_dollar_amount(amount):
    """Format dollar amount for LaTeX."""
    if amount == 0:
        return '\\$0'
    return '\\$' + format_number_with_commas(amount)

def main():
    excel_file = 'ComputerCrime_AUG2026_Impact_Final_With_C5C7_IFERROR_CHARTS_FIXED.xlsx'
    
    print('Loading workbook...')
    wb = load_workbook(excel_file, data_only=True)
    
    years = ['2016', '2017', '2018', '2019', '2020', '2021', '2022']
    
    # Extract data
    print('Extracting American Samoa data...')
    american_samoa_data = extract_extortion_data(wb, 'AMERICAN_SAMOA', years)
    
    print('Extracting California data...')
    california_data = extract_extortion_data(wb, 'CALIFORNIA', years)
    
    # Calculate means
    print('Calculating means...')
    american_samoa_means = calculate_means(american_samoa_data)
    california_means = calculate_means(california_data)
    
    # Determine distribution type
    dist_type, lower_bound, upper_bound = determine_distribution_type(
        american_samoa_means['mean_count'],
        california_means['mean_count']
    )
    
    # Prepare results dictionary
    results = {
        'american_samoa': {
            'raw_counts': american_samoa_data['raw_counts'],
            'raw_impacts': american_samoa_data['raw_impacts'],
            'scaled_counts': american_samoa_data['scaled_counts'],
            'scaled_impacts': american_samoa_data['scaled_impacts'],
            'mean_scaled_count': american_samoa_means['mean_count'],
            'mean_scaled_impact': american_samoa_means['mean_impact']
        },
        'california': {
            'raw_counts': california_data['raw_counts'],
            'raw_impacts': california_data['raw_impacts'],
            'scaled_counts': california_data['scaled_counts'],
            'scaled_impacts': california_data['scaled_impacts'],
            'mean_scaled_count': california_means['mean_count'],
            'mean_scaled_impact': california_means['mean_impact']
        },
        'distribution': {
            'type': dist_type,
            'lower_bound': lower_bound,
            'upper_bound': upper_bound,
            'range': upper_bound - lower_bound + 1 if dist_type != 'normal' else None
        },
        'years': years
    }
    
    # Print summary
    print('\n=== EXTRACTION SUMMARY ===')
    print(f'\nAmerican Samoa:')
    print(f'  Raw counts: {american_samoa_data["raw_counts"]}')
    print(f'  Raw impacts: {american_samoa_data["raw_impacts"]}')
    print(f'  Scaled counts: {american_samoa_data["scaled_counts"]}')
    print(f'  Scaled impacts: {american_samoa_data["scaled_impacts"]}')
    print(f'  Mean scaled count: {american_samoa_means["mean_count"]:.2f}')
    print(f'  Mean scaled impact: ${american_samoa_means["mean_impact"]:.2f}')
    
    print(f'\nCalifornia:')
    print(f'  Raw counts: {california_data["raw_counts"]}')
    print(f'  Raw impacts: {california_data["raw_impacts"]}')
    print(f'  Scaled counts: {california_data["scaled_counts"]}')
    print(f'  Scaled impacts: {california_data["scaled_impacts"]}')
    print(f'  Mean scaled count: {california_means["mean_count"]:.2f}')
    print(f'  Mean scaled impact: ${california_means["mean_impact"]:.2f}')
    
    print(f'\nDistribution:')
    print(f'  Type: {dist_type}')
    print(f'  Lower bound: {lower_bound}')
    print(f'  Upper bound: {upper_bound}')
    if dist_type != 'normal':
        print(f'  Range: {upper_bound - lower_bound + 1} values')
    
    # Save results to JSON
    with open('extortion_analysis_results.json', 'w') as f:
        json.dump(results, f, indent=2)
    
    print('\nResults saved to extortion_analysis_results.json')
    
    return results

if __name__ == '__main__':
    main()
