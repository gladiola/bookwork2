#!/usr/bin/env python3
"""
Final verification of the Extortion analysis project.
"""

import os
import json
from pathlib import Path

def check_file_exists(filepath, description):
    """Check if a file exists and return status."""
    if os.path.exists(filepath):
        size = os.path.getsize(filepath)
        print(f"✓ {description}: {filepath} ({size:,} bytes)")
        return True
    else:
        print(f"✗ {description}: {filepath} - NOT FOUND")
        return False

def main():
    print("=" * 80)
    print("EXTORTION ANALYSIS PROJECT - FINAL VERIFICATION")
    print("=" * 80)
    
    base_path = "/home/runner/work/bookwork2/bookwork2/KandRStyle"
    
    # Check data files
    print("\n### Data Files ###")
    data_files = [
        ("research/risk_air_gap/extortion_analysis_results.json", "Analysis results JSON"),
        ("excerpts/extortion_latex_output.txt", "LaTeX output with populated data"),
        ("excerpts/EXTORTION_ANALYSIS_SUMMARY.md", "Summary document"),
    ]
    
    all_data = all(check_file_exists(os.path.join(base_path, f), d) for f, d in data_files)
    
    # Check scripts
    print("\n### Python Scripts ###")
    scripts = [
        ("research/risk_air_gap/process_extortion_data.py", "Data extraction script"),
        ("research/risk_air_gap/generate_extortion_latex.py", "LaTeX generation script"),
        ("research/risk_air_gap/create_extortion_worksheet.py", "Excel worksheet creation"),
        ("research/risk_air_gap/generate_extortion_figures.py", "Figure generation script"),
        ("research/risk_air_gap/generate_exceedance_curves.py", "Exceedance curves generator"),
        ("research/risk_air_gap/update_workbook_toc.py", "TOC and hyperlink management"),
    ]
    
    all_scripts = all(check_file_exists(os.path.join(base_path, f), d) for f, d in scripts)
    
    # Check Excel workbook
    print("\n### Excel Workbook ###")
    workbook_path = os.path.join(base_path, "research/risk_air_gap/ComputerCrime_AUG2026_Impact_Final_With_C5C7_IFERROR_CHARTS_FIXED.xlsx")
    workbook_ok = check_file_exists(workbook_path, "Updated workbook")
    
    if workbook_ok:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(workbook_path, read_only=True)
            print(f"  - Total worksheets: {len(wb.sheetnames)}")
            if 'Discrete_Extortion' in wb.sheetnames:
                print(f"  ✓ Discrete_Extortion worksheet found")
            else:
                print(f"  ✗ Discrete_Extortion worksheet NOT found")
            wb.close()
        except Exception as e:
            print(f"  ! Error loading workbook: {e}")
    
    # Check figures
    print("\n### PNG Figures (Black and White) ###")
    figures = [
        ("figures/discrete_extortion_count.png", "Extortion discrete count chart"),
        ("figures/histogram_extortion_impact.png", "Extortion impact histogram"),
        ("figures/percent_extortion_impact.png", "Extortion impact exceedance curve"),
        ("figures/histogram_phishing_impact.png", "Phishing impact histogram (updated)"),
        ("figures/percent_phishing_impact.png", "Phishing exceedance curve (updated)"),
        ("figures/histogram_threats_impact.png", "Threats impact histogram (updated)"),
        ("figures/percent_threats_impact.png", "Threats exceedance curve (updated)"),
        ("figures/histogram_idtheft_impact.png", "ID Theft impact histogram (updated)"),
        ("figures/percent_idtheft_impact.png", "ID Theft exceedance curve (updated)"),
    ]
    
    all_figures = all(check_file_exists(os.path.join(base_path, f), d) for f, d in figures)
    
    # Load and display key results
    print("\n### Key Results ###")
    results_path = os.path.join(base_path, "research/risk_air_gap/extortion_analysis_results.json")
    if os.path.exists(results_path):
        with open(results_path, 'r') as f:
            results = json.load(f)
        
        print(f"\nDistribution Type: {results['distribution']['type'].upper()}")
        print(f"Lower Bound: {results['distribution']['lower_bound']} events")
        print(f"Upper Bound: {results['distribution']['upper_bound']} events")
        print(f"Range: {results['distribution']['range']} discrete values")
        
        if 'monte_carlo' in results and 'impact_percentiles' in results['monte_carlo']:
            p = results['monte_carlo']['impact_percentiles']
            print(f"\nImpact Percentiles:")
            print(f"  P10 (90% probability): ${p['p10']:,.2f}")
            print(f"  P50 (50% probability): ${p['p50']:,.2f}")
            print(f"  P90 (10% probability): ${p['p90']:,.2f}")
            print(f"  80% CI: [${p['ci_lower']:,.2f}, ${p['ci_upper']:,.2f}]")
    
    # Final summary
    print("\n" + "=" * 80)
    print("VERIFICATION SUMMARY")
    print("=" * 80)
    
    status = []
    status.append(("Data files", all_data))
    status.append(("Python scripts", all_scripts))
    status.append(("Excel workbook", workbook_ok))
    status.append(("PNG figures", all_figures))
    
    all_ok = all(ok for _, ok in status)
    
    for item, ok in status:
        symbol = "✓" if ok else "✗"
        print(f"{symbol} {item}: {'OK' if ok else 'ISSUES FOUND'}")
    
    print("\n" + "=" * 80)
    if all_ok:
        print("✓✓✓ ALL VERIFICATIONS PASSED ✓✓✓")
    else:
        print("✗✗✗ SOME VERIFICATIONS FAILED ✗✗✗")
    print("=" * 80)
    
    return 0 if all_ok else 1

if __name__ == '__main__':
    exit(main())
